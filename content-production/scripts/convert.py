"""
Convert binary documents to structured Markdown — preserves layout.

Companion to ``extract``: when you want *formatting* (not just raw words), use
convert.  When you want it *fast* and raw, use extract.

Supported this round:
    - DOCX → Markdown (headings, lists, bold/italic, tables)

DOCX path:
    1. Primary: ``python-docx`` walks the document in order and maps paragraph
       styles + run formatting to Markdown (precise style mapping).
    2. Fallback: when the DOCX XML is non-compliant and python-docx cannot load
       it, ``mammoth`` converts the file leniently as a safety net.

Images and tables are extracted as sidecar assets (restored from the original
solo-skills ``content-layout-expert``), not left as text placeholders:

    - Images: each inline/anchored picture is written to the media directory as
      ``{stem}-image-{n}.{ext}`` and emitted as its own Markdown block
      ``![](path)``, fully separated from surrounding text. Images are resolved
      in document order via their drawing relationships, so file numbering
      always matches the reference order.
    - Tables: each table is written to ``{stem}-table-{n}.csv`` (clean data
      handoff for downstream rendering). The Markdown keeps a GitHub-flavoured
      table by default; pass ``tables_as_images`` to emit
      ``![](stem-table-{n}.png)`` instead — the WeChat pipeline, where the PNG
      is rendered later by ``media-composer``.

The Markdown is tidied for readability: ``·``/``•`` bullet paragraphs become
tight Markdown lists, and ``force_first_h1`` / ``normalize_headings`` fix the
title and section heading hierarchy that article DOCX files often carry
inconsistently.

PPTX is planned but not implemented yet.
"""

from __future__ import annotations

import csv
import io
import os
import posixpath
import re
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from scripts.common import logger

# Paragraphs whose text starts with one of these are treated as list items and
# grouped into tight Markdown lists (a common convention in CJK documents).
_BULLET_RE = re.compile(r"^[·•・‧]\s*")

# Heading-level heuristics for --normalize-headings: CJK numbered sections are
# mapped to a consistent H2/H3 hierarchy regardless of the source's (often
# inconsistent) heading styles. Tuned for the DOCX → WeChat article pipeline.
_CJK_NUM = "一二三四五六七八九十"
_HEADING_LEVEL_RULES = [
    (re.compile(rf"^[{_CJK_NUM}]+[、.．]"), 2),                    # 一、 二、 一.
    (re.compile(rf"^第[{_CJK_NUM}\d]+[章节部分篇]"), 2),            # 第一章 / 第1章
    (re.compile(r"^(?:结语|前言|序言|引言|总结|附录|后记|导言|引子)"), 2),
    (re.compile(r"^\d+[、.．)]"), 3),                              # 1. 1、 1)
    (re.compile(rf"^[(（][{_CJK_NUM}\d]+[)）]"), 3),              # (一) （1）
    (re.compile(r"^[①②③④⑤⑥⑦⑧⑨⑩⑪⑫]"), 3),                        # ① ②
]


def iter_block_items(parent):
    """Yield each Paragraph and Table child of *parent* in document order."""
    if isinstance(parent, _Cell):
        parent_elm = parent._tc
    elif isinstance(parent, DocumentType):
        parent_elm = parent.element.body
    else:
        raise ValueError("Unsupported parent type for iter_block_items")
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


class _ConvertState:
    """Mutable state threaded through one DOCX conversion.

    Carries the asset directory + reference prefix, the per-document image and
    table counters, and the lists of extracted asset references. Kept as a
    plain holder so the block-walking functions stay free of globals.
    """

    def __init__(
        self,
        media_dir: Path,
        ref_prefix: str,
        stem: str,
        extract_images: bool,
        extract_tables: bool,
        tables_as_images: bool,
        min_image_bytes: int,
        force_first_h1: bool,
        normalize_headings: bool,
    ) -> None:
        self.media_dir = media_dir
        self.ref_prefix = ref_prefix
        self.stem = stem
        self.extract_images = extract_images
        self.extract_tables = extract_tables
        self.tables_as_images = tables_as_images
        self.min_image_bytes = min_image_bytes
        self.force_first_h1 = force_first_h1
        self.normalize_headings = normalize_headings
        self.image_counter = 0
        self.table_counter = 0
        self.images: list[str] = []   # Markdown refs for extracted images
        self.tables: list[str] = []   # Markdown refs for extracted tables


def convert_to_md(
    input_path: str | Path,
    fmt: str | None = None,
    *,
    media_dir: str | Path | None = None,
    md_output_path: str | Path | None = None,
    extract_images: bool = True,
    extract_tables: bool = True,
    tables_as_images: bool = False,
    min_image_bytes: int = 0,
    force_first_h1: bool = False,
    normalize_headings: bool = False,
) -> str:
    """Convert a document to structured Markdown.

    Parameters
    ----------
    input_path : str | Path
        Path to the source document.
    fmt : str | None
        Force a format (``docx``).  Inferred from the file extension when omitted.
    media_dir : str | Path | None
        Directory for extracted image and table assets.  Defaults to an
        ``assets`` subdirectory beside the source document (created if missing).
    md_output_path : str | Path | None
        Where the Markdown will be written.  Used to make image references
        relative to the Markdown file.
    extract_images : bool
        Extract embedded images to files and reference them inline (default).
        When False, images collapse back to ``(image)`` text placeholders.
    extract_tables : bool
        Extract tables to ``{stem}-table-{n}.csv`` sidecar files.
    tables_as_images : bool
        Reference tables as ``![]({stem}-table-{n}.png)`` instead of a
        GitHub-flavoured Markdown table — for the WeChat pipeline, where the
        PNG is rendered from the CSV by ``media-composer``.
    min_image_bytes : int
        Skip images smaller than this many bytes (decorative bullets/icons).
        ``0`` keeps every image (default).
    force_first_h1 : bool
        Force the first prose paragraph to an H1 title (``# Title``), stripping
        bold/italic markers. Mirrors how article DOCX files carry the title as
        a bold paragraph rather than a Heading style.
    normalize_headings : bool
        Remap heading levels to a clean hierarchy using CJK numbering patterns
        (``一、`` → H2, ``1.`` → H3, ``结语`` → H2), ignoring the source's often
        inconsistent heading styles. Combine with ``force_first_h1`` for the
        full ``title=H1 / sections=H2 / subsections=H3`` structure.

    Returns
    -------
    str
        The converted Markdown text.
    """
    path = Path(input_path)
    if not path.exists():
        raise FileNotFoundError(f"Input file not found: {path}")

    fmt = (fmt or path.suffix.lstrip(".")).lower()
    if fmt == "docx":
        media = _resolve_media_dir(media_dir, path)
        media.mkdir(parents=True, exist_ok=True)
        ref_prefix = _ref_prefix(media, md_output_path)
        state = _ConvertState(
            media, ref_prefix, path.stem,
            extract_images, extract_tables, tables_as_images,
            min_image_bytes, force_first_h1, normalize_headings,
        )
        markdown = _convert_docx(path, state)
        _log_assets(path.name, state)
        return markdown
    if fmt == "pptx":
        raise NotImplementedError(
            "convert for .pptx is planned but not yet implemented "
            "(this round ships docx)."
        )
    raise ValueError(
        f"Unsupported format for convert: '.{fmt}' (supported: docx)"
    )


def _resolve_media_dir(
    media_dir: str | Path | None,
    input_path: Path,
) -> Path:
    """Pick the asset directory: explicit flag → ``<docx dir>/assets``."""
    if media_dir:
        return Path(media_dir)
    return input_path.parent / "assets"


def _ref_prefix(media_dir: Path, md_output_path: str | Path | None) -> str:
    """Path prefix (with trailing ``/``) from the Markdown file to the assets.

    Returns ``""`` when assets sit next to the Markdown so references use just
    the filename. Paths use ``/`` regardless of OS so they work in Markdown.
    """
    start = Path(md_output_path).resolve().parent if md_output_path else Path.cwd()
    try:
        rel = os.path.relpath(media_dir.resolve(), start)
    except ValueError:  # different drive letters (Windows) — fall back to literal
        rel = str(media_dir)
    rel = rel.replace(os.sep, "/")
    if rel in (".", ""):
        return ""
    return rel.rstrip("/") + "/"


# ---------------------------------------------------------------------------
# DOCX → Markdown (python-docx primary, mammoth fallback)
# ---------------------------------------------------------------------------


_RELS_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def _dangling_rel_ids(zf: zipfile.ZipFile, rels_name: str) -> set[str]:
    """Return relationship ids in *rels_name* whose target is missing from the zip.

    Word sometimes writes a relationship pointing at a target that was never
    stored (e.g. ``Target="../NULL"`` for an image that failed to embed).
    python-docx eagerly loads every part when opening the package, so a single
    dangling entry makes the whole file unreadable.
    """
    try:
        root = ET.fromstring(zf.read(rels_name))
    except (KeyError, ET.ParseError):
        return set()

    names = set(zf.namelist())
    # Targets resolve against the directory holding the part that owns the
    # .rels file, i.e. the parent of the "_rels" folder ("" for the root
    # "_rels/.rels", "word" for "word/_rels/document.xml.rels").
    base = posixpath.dirname(posixpath.dirname(rels_name))
    dangling: set[str] = set()
    for rel in root.findall(f"{{{_RELS_NS}}}Relationship"):
        target = rel.get("Target") or ""
        if rel.get("TargetMode") == "External" or not target or "://" in target:
            continue
        resolved = posixpath.normpath(posixpath.join(base, target)).lstrip("/")
        if resolved not in names:
            dangling.add(rel.get("Id") or "")
            logger.warning(
                "dropping dangling relationship %s -> %s in %s",
                rel.get("Id"), target, rels_name,
            )
    return {rid for rid in dangling if rid}


def _repair_docx_bytes(path: Path) -> io.BytesIO | None:
    """Rebuild *path* in memory with dangling relationships removed.

    Returns ``None`` when nothing was repairable, so the caller can fall
    through to the lenient converter. The file on disk is never modified.
    """
    try:
        with zipfile.ZipFile(path) as zf:
            rels_names = [n for n in zf.namelist() if n.endswith(".rels")]
            repairs = {n: _dangling_rel_ids(zf, n) for n in rels_names}
            repairs = {n: ids for n, ids in repairs.items() if ids}
            if not repairs:
                return None

            buffer = io.BytesIO()
            with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as out:
                for item in zf.infolist():
                    data = zf.read(item.filename)
                    drop = repairs.get(item.filename)
                    if drop:
                        data = _strip_relationships(data, drop)
                    out.writestr(item, data)
            buffer.seek(0)
            return buffer
    except (zipfile.BadZipFile, OSError) as exc:
        logger.warning("could not repair %s (%s)", path.name, exc)
        return None


def _strip_relationships(data: bytes, drop_ids: set[str]) -> bytes:
    """Remove the ``Relationship`` entries named in *drop_ids* from a .rels part."""
    try:
        root = ET.fromstring(data)
    except ET.ParseError:
        return data
    for rel in list(root.findall(f"{{{_RELS_NS}}}Relationship")):
        if (rel.get("Id") or "") in drop_ids:
            root.remove(rel)
    ET.register_namespace("", _RELS_NS)
    return ET.tostring(root, encoding="UTF-8", xml_declaration=True)


def _load_document(path: Path) -> DocumentType | None:
    """Open *path* with python-docx, repairing dangling relationships if needed."""
    try:
        return Document(str(path))
    except Exception as exc:
        repaired = _repair_docx_bytes(path)
        if repaired is None:
            logger.warning("python-docx could not load %s (%s)", path.name, exc)
            return None
        try:
            doc = Document(repaired)
        except Exception as repair_exc:
            logger.warning(
                "python-docx could not load %s even after repair (%s)",
                path.name, repair_exc,
            )
            return None
        logger.info("loaded %s after dropping dangling relationships", path.name)
        return doc


def _convert_docx(path: Path, state: _ConvertState) -> str:
    doc = _load_document(path)
    if doc is None:  # non-compliant / corrupted XML
        logger.warning(
            "falling back to mammoth for %s (image/table extraction skipped)",
            path.name,
        )
        return _convert_docx_mammoth(path)

    lines: list[str] = []
    # Track numbered-list numbering across consecutive numbered paragraphs.
    list_counter = 0
    in_numbered_list = False
    first_heading_done = False

    for block in iter_block_items(doc):
        if isinstance(block, Paragraph):
            body, list_type, image_refs = _paragraph_to_md(block, state)
            if body is None and not image_refs:
                # An empty paragraph breaks a running numbered list.
                in_numbered_list = False
                continue
            # Force the first prose paragraph to an H1 title when requested.
            if (state.force_first_h1 and not first_heading_done
                    and body is not None and list_type is None):
                body, became_title = _force_h1(body)
                if became_title:
                    first_heading_done = True
            # Emit the text body (with any list prefix).
            if body is not None:
                if list_type == "number":
                    if not in_numbered_list:
                        list_counter = 0
                        in_numbered_list = True
                    list_counter += 1
                    lines.append(f"{list_counter}. {body}")
                else:
                    in_numbered_list = False
                    if list_type == "bullet":
                        lines.append(f"- {body}")
                    else:
                        lines.append(body)
            else:
                in_numbered_list = False
            # Emit images as their own blocks, fully separated from the text.
            for ref in image_refs:
                lines.append(ref)
                in_numbered_list = False
        elif isinstance(block, Table):
            in_numbered_list = False
            table_md = _table_to_md(block, state)
            if table_md:
                lines.append(table_md)

    lines = _group_bullets(lines)
    logger.info(
        "DOCX convert: %d blocks from %s", len(lines), path.name,
    )
    return "\n\n".join(lines)


def _convert_docx_mammoth(path: Path) -> str:
    """Lenient fallback: mammoth's built-in DOCX → Markdown conversion."""
    import mammoth

    path = Path(path)
    with open(path, "rb") as fh:
        result = mammoth.convert_to_markdown(fh)
    logger.info("DOCX convert via mammoth fallback: %s", path.name)
    return result.value


def _heading_level(paragraph: Paragraph) -> int | None:
    """Return Markdown heading level (1–6) for the paragraph, or None.

    Checks the paragraph style name first ("Heading N", "Title"), then falls
    back to the outline level in the XML (which survives custom style names).
    """
    style_name = (paragraph.style.name if paragraph.style else "") or ""
    low = style_name.lower()
    if low.startswith("heading") or low.startswith("标题"):
        m = re.search(r"(\d+)", style_name)
        n = int(m.group(1)) if m else 1
        return max(1, min(n, 6))
    if low == "title" or low == "标题":
        return 1

    pPr = paragraph._element.find(qn("w:pPr"))
    if pPr is not None:
        outline = pPr.find(qn("w:outlineLvl"))
        if outline is not None:
            val = outline.get(qn("w:val"))
            if val is not None and val.lstrip("-").isdigit():
                n = int(val) + 1  # outlineLvl is 0-based
                if 1 <= n <= 6:
                    return n
    return None


def _list_kind(paragraph: Paragraph) -> str | None:
    """Return ``"bullet"`` / ``"number"`` for list-style paragraphs, else None."""
    name = (paragraph.style.name if paragraph.style else "") or ""
    low = name.lower()
    if "bullet" in low or "项目符号" in low:
        return "bullet"
    if "number" in low or "编号" in low:
        return "number"
    return None


def _count_images(paragraph: Paragraph) -> int:
    """Count inline images (drawings) carried by the paragraph's runs.

    Used only for the lightweight ``(image)`` placeholder fallback when image
    extraction is disabled.
    """
    count = 0
    for run in paragraph.runs:
        if "pic:pic" in run._element.xml or "<w:drawing" in run._element.xml:
            count += 1
    return count


def _guess_extension(data: bytes) -> str:
    """Guess an image extension from magic bytes (fallback only)."""
    if data[:8] == b"\x89PNG\r\n\x1a\n":
        return ".png"
    if data[:2] == b"\xff\xd8":
        return ".jpg"
    if data[:6] in (b"GIF87a", b"GIF89a"):
        return ".gif"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return ".webp"
    if data[:4] == b"BM":
        return ".bmp"
    return ".png"


def _image_extension(part, blob: bytes) -> str:
    """Image extension from the package part name, falling back to magic bytes."""
    partname = getattr(part, "partname", None)
    ext = getattr(partname, "ext", None)
    if ext:
        return f".{ext.lstrip('.').lower()}"
    return _guess_extension(blob)


def _resolve_image_refs(paragraph: Paragraph, state: _ConvertState) -> list[str]:
    """Resolve inline/anchored images in *paragraph* to extracted files.

    Walks each ``a:blip`` in document order, resolves its relationship to the
    image part, writes the bytes to the media directory as
    ``{stem}-image-{n}.{ext}``, and returns the Markdown image references that
    replace the old ``(image)`` placeholders. Numbering follows document order,
    so references and files always align.
    """
    refs: list[str] = []
    for blip in paragraph._element.iter(qn("a:blip")):
        rId = blip.get(qn("r:embed")) or blip.get(qn("r:link"))
        if not rId:
            continue
        part = paragraph.part.related_parts.get(rId)
        if part is None:
            logger.warning("image relationship %s not found — skipping", rId)
            continue
        blob = getattr(part, "blob", b"") or b""
        if len(blob) < state.min_image_bytes:
            logger.info(
                "skipping image %s (%d bytes < %d)", rId, len(blob), state.min_image_bytes,
            )
            continue

        state.image_counter += 1
        n = state.image_counter
        filename = f"{state.stem}-image-{n}{_image_extension(part, blob)}"
        (state.media_dir / filename).write_bytes(blob)

        ref = f"{state.ref_prefix}{filename}"
        state.images.append(ref)
        refs.append(f"![{state.stem}-image-{n}]({ref})")
    return refs


def _format_run(run) -> str:
    """Render a single run's text with bold/italic Markdown markers."""
    text = run.text or ""
    if not text:
        return ""
    bold = bool(run.bold)
    italic = bool(run.italic)
    if bold and italic:
        return f"***{text}***"
    if bold:
        return f"**{text}**"
    if italic:
        return f"*{text}*"
    return text


def _force_h1(body: str) -> tuple[str, bool]:
    """Render *body* (text only) as an H1 title line.

    Strips existing heading markers and bold/italic markers for a clean title.
    Returns ``(rendered, became_title)``; ``became_title`` is False when *body*
    has no text after stripping, so the caller leaves the title slot open for a
    later paragraph. Image references are emitted separately by the caller.
    """
    text = re.sub(r"^\s*#+\s*", "", body)    # drop existing heading marks
    text = re.sub(r"\*+", "", text).strip()  # drop bold/italic markers
    if not text:
        return body, False
    return f"# {text}", True


def _normalize_heading_level(text: str, detected: int, reserve_h1: bool) -> int:
    """Pick a heading level for *text*, ignoring the source's (often messy) styles.

    CJK numbered sections map to a fixed hierarchy (``一、`` → 2, ``1.`` → 3,
    ``结语`` → 2) so sibling sections authored at different source levels end up
    consistent. Anything without a recognisable prefix keeps its detected level,
    demoted off H1 when *reserve_h1* is set (the title owns H1).
    """
    for pattern, level in _HEADING_LEVEL_RULES:
        if pattern.match(text):
            return level
    if reserve_h1 and detected <= 1:
        return 2
    return max(1, min(detected, 6))


def _paragraph_to_md(
    paragraph: Paragraph, state: _ConvertState,
) -> tuple[str | None, str | None, list[str]]:
    """Convert a paragraph to Markdown, splitting text from images.

    Returns ``(body, list_kind, image_refs)`` where:
      - ``body`` is the rendered text WITHOUT any list prefix (the caller adds
        ``- `` or ``N. ``), a heading line, or ``None`` for image-only / empty.
      - ``list_kind`` is ``"bullet"`` / ``"number"`` / ``None``.
      - ``image_refs`` is the list of ``![](path)`` references for the paragraph;
        the caller emits each as its own block so images never cling to text.
    """
    if state.extract_images:
        image_refs = _resolve_image_refs(paragraph, state)
    else:
        image_refs = ["(image)"] * _count_images(paragraph)

    body = "".join(_format_run(run) for run in paragraph.runs).strip()

    level = _heading_level(paragraph)
    if level is not None and body:
        clean = re.sub(r"\*+", "", body).strip()
        if state.normalize_headings:
            level = _normalize_heading_level(clean, level, state.force_first_h1)
        return f"{'#' * level} {clean}", None, image_refs

    if not body:
        return None, None, image_refs
    return body, _list_kind(paragraph), image_refs


def _table_rows(table: Table) -> list[list[str]]:
    """Read a DOCX table into a list of rows of stripped cell text."""
    rows: list[list[str]] = []
    for row in table.rows:
        rows.append([cell.text.strip() for cell in row.cells])
    return rows


def _write_csv(path: Path, rows: list[list[str]]) -> None:
    """Write *rows* to a CSV (UTF-8 with BOM, so Excel opens CJK correctly)."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh)
        writer.writerows(rows)


def _rows_to_gfm(rows: list[list[str]]) -> str:
    """Render rows as a GitHub-flavoured Markdown table."""
    md_rows = []
    for row in rows:
        cells = [
            cell.replace("|", "\\|").replace("\n", " ") for cell in row
        ]
        md_rows.append("| " + " | ".join(cells) + " |")
    ncols = max(len(rows[0]), 1)
    separator = "| " + " | ".join(["---"] * ncols) + " |"
    return "\n".join([md_rows[0], separator, *md_rows[1:]])


def _table_to_md(table: Table, state: _ConvertState) -> str:
    """Render a DOCX table, extracting a CSV sidecar for downstream rendering.

    The Markdown representation depends on ``state.tables_as_images``: a
    GitHub-flavoured table (default, readable everywhere) or an image reference
    ``![]({stem}-table-{n}.png)`` for the WeChat pipeline — the PNG is rendered
    from the extracted CSV by ``media-composer``. The CSV is always written
    when table extraction is on, regardless of the Markdown representation.
    """
    rows = _table_rows(table)
    if not rows:
        return ""

    if not state.extract_tables:
        return _rows_to_gfm(rows)

    state.table_counter += 1
    n = state.table_counter
    csv_name = f"{state.stem}-table-{n}.csv"
    _write_csv(state.media_dir / csv_name, rows)
    state.tables.append(f"{state.ref_prefix}{csv_name}")

    if state.tables_as_images:
        png_ref = f"{state.ref_prefix}{state.stem}-table-{n}.png"
        return f"![{state.stem}-table-{n}]({png_ref})"
    return _rows_to_gfm(rows)


def _group_bullets(lines: list[str]) -> list[str]:
    """Convert ``·``/``•``-prefixed paragraphs into tight Markdown lists.

    Consecutive bullet blocks are joined into one element (items separated by a
    single newline) so they render as a compact list rather than loose blocks.
    """
    grouped: list[str] = []
    i = 0
    while i < len(lines):
        if _BULLET_RE.match(lines[i]):
            items: list[str] = []
            while i < len(lines) and _BULLET_RE.match(lines[i]):
                items.append(_BULLET_RE.sub("- ", lines[i]))
                i += 1
            grouped.append("\n".join(items))
        else:
            grouped.append(lines[i])
            i += 1
    return grouped


def _log_assets(name: str, state: _ConvertState) -> None:
    """Summarise extracted assets to the log (nothing when the doc had none)."""
    if state.images or state.tables:
        logger.info(
            "convert assets → %s: %d image(s), %d table(s)",
            state.media_dir, len(state.images), len(state.tables),
        )
